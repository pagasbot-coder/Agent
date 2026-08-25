#!/usr/bin/env python3
"""Сборка шпаргалки PM на собес (формат как в видеo 2rfo3Q9ie1c) → .xlsx для Google Таблиц."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "interview-pm-cheat-sheet.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header_row(ws, row: int, cols: int):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = WRAP


def set_widths(ws, widths: dict[int, float]):
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def sheet_nav(wb: Workbook):
    ws = wb.active
    ws.title = "Навигация"
    rows = [
        ("Лист", "Зачем", "Когда открывать"),
        ("Легенда", "Самопрезентация 60–90 сек + блоки карьеры", "Старт собеса, «расскажите о себе»"),
        ("Кейсы", "Ответы на типовые вопросы PM", "Тех/поведенческое интервью"),
        ("Exolve", "Доп. блок под МТС Exolve / антифрод", "Если вакансия compliance / telecom"),
        ("ХардСофт", "Hard + soft skills", "«Чем сильнее слабее»"),
        ("Словарик", "Термины одной строкой", "Telecom / compliance"),
        ("Вопросы", "Что спросить у них", "«Ваши вопросы?»"),
        ("Посмотреть", "Ссылки и напоминания", "Подготовка дома"),
    ]
    for r, row in enumerate(rows, 1):
        ws.append(row)
    style_header_row(ws, 1, 3)
    set_widths(ws, {1: 14, 2: 42, 3: 38})
    ws.freeze_panes = "A2"


def sheet_legend(wb: Workbook):
    ws = wb.create_sheet("Легенда")
    ws.merge_cells("A1:F1")
    ws["A1"] = "Самопрезентация"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = [
        "До IT / почему PM",
        "Какой проект (последний / целевой)",
        "Состав команды (типовой)",
        "За что отвечал",
        "Почему ищешь новое",
        "Что ищешь сейчас (из вакансии)",
    ]
    answers = [
        "Госконтур (КБПК, администрация Кириш), ритейл («Твоя цена», контроль/СБ), потом IT: телеком, EdTech, PM-инструменты. Перешёл в PM ради измеримого результата и управления сроками/рисками.",
        "Целевой: МТС Exolve — compliance антифрод + ГИС «Антифрод». Параллельно: ProductMap (EdTech, PM frameworks). Beeline: интеграции, Problem Management, observability доступности с 23% до 97%.",
        "PM + продукт + юристы + ИБ + DevOps + разработка + сеть/телефония + support. На Exolve — уточнить на собесе фактический состав.",
        "Сроки под регуляторику и НПА, реестр обязательств, gap по продуктам, риски, релизы, статус стейкхолдерам, интеграции, runbook до go-live.",
        "[ЗАПОЛНИ под вакансию: 1 факт + 1 цель, без негатива]",
        "Проект с жёсткими дедлайнами закона, где PM ведёт план-график и метрики соответствия; влияние на бизнес (FP, SLA ГИС), не «ML ради ML».",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(2, i, h)
    for i, a in enumerate(answers, 1):
        ws.cell(3, i, a)
    style_header_row(ws, 2, 6)

    ws["A5"] = "Intro 60–90 сек (выучить)"
    ws["A5"].font = BOLD
    ws.merge_cells("A5:F5")
    intro = (
        "PM с опытом интеграций, рисков и дедлайнов в телекоме и госконтуре. "
        "На Exolve вижу проект compliance + ГИС «Антифрод»: реестр НПА → gap по продуктам → "
        "план под даты регулятора → интеграции и ops → метрики ложных блокировок для B2B. "
        "СОРМ и антифрод не смешиваю. Прямых переговоров с РКН и своего antifraud engine не строил — "
        "близко Beeline (интеграции, сроки, observability 23→97%) и работа с НПА в госконтуре."
    )
    ws.merge_cells("A6:F8")
    ws["A6"] = intro
    ws["A6"].alignment = WRAP

    for row in ws.iter_rows(min_row=2, max_row=8, min_col=1, max_col=6):
        for cell in row:
            cell.alignment = WRAP
    set_widths(ws, {i: 28 for i in range(1, 7)})


def sheet_cases(wb: Workbook):
    ws = wb.create_sheet("Кейсы")
    ws.append(["Вопрос", "Ответ"])
    style_header_row(ws, 1, 2)
    cases = [
        (
            "Нет документации на проекте",
            "Привлечь системного аналитика или бывшего сотрудника для интервью. "
            "Параллельно правило: новая фича/интеграция → артефакт в wiki до релиза. "
            "Кontрольные точки у PM: что, кто, когда. Не завязать знания на одном человеке.",
        ),
        (
            "Уходит главный / ключевой разработчик",
            "1-on-1: сценарий «остаётся / уходит». Если уходит — описание архитектуры, "
            "handover 1–1,5 мес, сразу найм, подключить соседний юнит. "
            "Профилактика: bus factor в реестре рисков с первого месяца.",
        ),
        (
            "Техдолг — как приоритизируешь",
            "Блокер релиза > влияние на KPI/фичи > быстрые wins. "
            "Окно между релизами или конец года. Compliance must-have — отдельная очередь, не «удобный рефакторинг».",
        ),
        (
            "Инцидент в пятницу вечером",
            "Сначала стабилизация: runbook, дежурный, эскалация. Потом retro — почему и как не повторить. "
            "Для ГИС отдельно: SLA 24/7, RTO/RPO (уточнить у Exolve).",
        ),
        (
            "Inhouse vs аутстаф / аутсорс",
            "Inhouse: погружение, скорость изменений, дороже, хуже масштаб. "
            "Аутстаф: гибко и дешевле, нужны контрольные точки. "
            "Ответственность перед регулятором не делегирую «на подрядчика без RACI».",
        ),
        (
            "Проблема, которую никто не может решить",
            "Декомпозиция → аналоги → эксперты внутри/снаружи → гипотезы → решение. "
            "С бизнесом — язык денег: критичность vs стоимость поиска (как с диагностикой авто).",
        ),
        (
            "Была документация на проекте?",
            "Устав, ТЗ/ФТ, user stories, реестр рисков, runbook. "
            "На compliance — реестр НПА как single source of truth. PM/Legal следят за актуальностью.",
        ),
        (
            "Онбординг разработчика без документации",
            "Intro → доступы → чаты → buddy → мелкие задачи → фиксировать находки в wiki с day 1. "
            "Параллельно процесс обязательной документации.",
        ),
        (
            "Спор разработчиков (стек / Python vs Go)",
            "Синки 1-on-1 → общий митинг, медиатор. Язык цифр: скорость MVP, поддержка, масштаб. "
            "Решение в ADR. Для MVP — что быстрее к ценности; для платформы — что жить 3+ года.",
        ),
        (
            "От кого задачи / стейкхолдеры",
            "Спонсор, продукт, юристы, ИБ, CTO/архитектор, ops, support. "
            "На Exolve: кто A по сроку регулятора — имя, не «отдел».",
        ),
        (
            "Прямые подчинённые?",
            "Честно: в Scrum часто без formal line, но delivery, приоритеты, риски, статус — на PM. "
            "RACI вместо «все подчинены мне».",
        ),
        (
            "Когда документируешь задачу",
            "Discovery → Delivery: user story, критерии приёмки, задачи. "
            "Системная аналитика/API — до кода. Confluence — до merge для новых интеграций.",
        ),
    ]
    for q, a in cases:
        ws.append([q, a])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = WRAP
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 72
    ws.freeze_panes = "A2"


def sheet_exolve(wb: Workbook):
    ws = wb.create_sheet("Exolve")
    ws.append(["Тема", "Ответ / факт"])
    style_header_row(ws, 1, 2)
    rows = [
        ("Суть Exolve", "Голос, SMS, API, 8-800, B2B. Антифрод = compliance + продукт + бизнес."),
        ("СОРМ ≠ антифрод", "СОРМ: 126+144 ФЗ, доступ органов. Антифрод: ГИС, сигналы, блокировки, маркировка."),
        ("Подход PM", "Реестр НПА → gap matrix → план под даты → интеграции + ops → FP rate для B2B."),
        ("90 дней (каркас)", "1–30: реестр, gap, RACI. 31–60: MVP интеграции, пилот. 61–90: scope, автоматизация, аудит."),
        ("Честные пробелы", "Нет прямых переговоров с РКН; нет deep СОРМ; не строил antifraud engine оператора."),
        ("Не обещать", "Даты НПА из СМИ; «уже построил ГИС»; смешение СОРМ и антифрода в одном эпике."),
        ("KPI проекта", "Coverage к дате N; 0 критичных предписаний; SLA ГИС; FP ≤ X%; time-to-block/report."),
    ]
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = WRAP
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 78
    ws.freeze_panes = "A2"


def sheet_hardsoft(wb: Workbook):
    ws = wb.create_sheet("ХардСофт")
    ws.append(["Категория", "Содержание"])
    style_header_row(ws, 1, 2)
    rows = [
        ("Hard — инструменты", "Jira, Confluence, Miro/Notion, статус-шаблон, диаграмма Ганта, реестр рисков."),
        ("Hard — техника (уровень PM)", "API/BFF, интеграции, observability, базово Docker/PostgreSQL. Не притворяться архитектором СОРМ."),
        ("Hard — сильная сторона", "Beeline: observability доступности 23% → 97%; интеграции; Problem Management."),
        ("Soft", "Медиатор, язык цифр, честные пробелы, 1-on-1, статус без сюрпризов, работа с юристами."),
        ("Anti-patterns", "Блеф по НПА; обещания без юриста; scope creep без change control."),
    ]
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = WRAP
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 78


def sheet_glossary(wb: Workbook):
    ws = wb.create_sheet("Словарик")
    ws.append(["Термин", "Одна строка"])
    style_header_row(ws, 1, 2)
    terms = [
        ("ГИС «Антифрод»", "Госсистема обмена сигналами; не путать с СОРМ"),
        ("СОРМ", "126+144 ФЗ — доступ уполномоченных органов"),
        ("FP / FN", "Ложная блокировка / пропуск фрода"),
        ("ВАТС", "Облачная телефония — отдельный legal + product impact"),
        ("Gap matrix", "Требование → продукт Exolve → статус"),
        ("Must / should", "Из реестра юриста, не из головы PM"),
        ("RACI", "R исполнитель, A ответственный за результат, C консультант, I в курсе"),
        ("Runbook", "Ops после go-live: дежурства, эскалации, RCA"),
        ("Time-to-block", "Скорость блокировки по сигналу"),
        ("Rolling wave", "Детальный план на 2–3 недели, дальше — полосы"),
    ]
    for t in terms:
        ws.append(list(t))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = WRAP
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 65


def sheet_questions(wb: Workbook):
    ws = wb.create_sheet("Вопросы")
    ws.append(["#", "Вопрос интервьюеру"])
    style_header_row(ws, 1, 2)
    qs = [
        "Какая роль Exolve: оператор / агрегатор / партнёр МТС?",
        "Какие продукты in scope первой волны: API, SMS, голос, 8-800, ВАТС?",
        "Кто accountable за дедлайн регулятора — конкретное имя?",
        "Узел обмена с ГИС: SLA 24/7, RTO/RPO?",
        "Есть ли бюджет false positive для B2B (порог X%)?",
        "Процесс апелляций / whitelist для легальных клиентов?",
        "Что уже сделано vs gap на старте?",
        "Как устроена связка PM — Legal — ИБ — DevOps?",
    ]
    for i, q in enumerate(qs, 1):
        ws.append([i, q])
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 75


def sheet_watch(wb: Workbook):
    ws = wb.create_sheet("Посмотреть")
    ws.append(["Что", "Ссылка / заметка"])
    style_header_row(ws, 1, 2)
    rows = [
        ("Видео-референс шпаргалки", "https://www.youtube.com/watch?v=2rfo3Q9ie1c"),
        ("Exolve prep (канон)", "quiet-partner/knowledge-base/mts-exolve-antifraud-project-prep.md"),
        ("Чеклист собеса", "quiet-partner/docs/mts-exolve-next-checklist.md"),
        ("Lifecycle playbook", "quiet-partner/knowledge-base/project-lifecycle-playbook.md"),
        ("Перед собесом", "Проговорить intro 60–90 сек без бумаги; СОРМ ≠ антифрод одной фразой"),
    ]
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 65


def main():
    wb = Workbook()
    sheet_nav(wb)
    sheet_legend(wb)
    sheet_cases(wb)
    sheet_exolve(wb)
    sheet_hardsoft(wb)
    sheet_glossary(wb)
    sheet_questions(wb)
    sheet_watch(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
